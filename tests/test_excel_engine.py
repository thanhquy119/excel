from datetime import datetime
from inspect import signature
from io import BytesIO

from openpyxl import Workbook, load_workbook
from streamlit.testing.v1 import AppTest

from excel_ai.models import Action, NormalizationOptions, OperationPlan
from excel_ai.operations import execute_plan, normalize_value
from excel_ai.planner import create_plan
from excel_ai.schema import inspect_workbook, schema_for_ai


def workbook_bytes(title: str, headers: list[str], rows: list[list[object]]) -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = title
    sheet.append(headers)
    for row in rows:
        sheet.append(row)
    buffer = BytesIO()
    workbook.save(buffer)
    workbook.close()
    return buffer.getvalue()


def test_app_starts_without_upload():
    app = AppTest.from_file("app.py", default_timeout=15).run()
    assert not app.exception
    assert len(app.file_uploader) == 1
    assert app.button[0].label == "Xử lý file"


def test_default_model_is_gemini_36_flash():
    assert signature(create_plan).parameters["model"].default == "gemini-3.6-flash"


def test_schema_sent_to_ai_contains_no_cell_values():
    data = workbook_bytes("HoaDon", ["Mã số thuế", "Số tiền"], [["0312345678", 150000]])
    schema = inspect_workbook(data, "hoa_don.xlsx", "F1")
    payload = str(schema_for_ai([schema]))

    assert "Mã số thuế" in payload
    assert "0312345678" not in payload
    assert "150000" not in payload


def test_compare_columns_highlights_matching_source_cells():
    source_data = workbook_bytes("BanRa", ["MST"], [["031-234-5678"], ["0109999999"], [None]])
    compare_data = workbook_bytes("KeToan", ["MST khách hàng"], [["0312345678"], ["0201111111"]])
    source_schema = inspect_workbook(source_data, "ban_ra.xlsx", "F1")
    compare_schema = inspect_workbook(compare_data, "ke_toan.xlsx", "F2")

    plan = OperationPlan(
        action=Action.COMPARE_COLUMNS,
        source_column_id=source_schema.sheets[0].columns[0].id,
        comparison_column_id=compare_schema.sheets[0].columns[0].id,
        highlight_condition="matched",
        color="FFF59D",
        normalization=NormalizationOptions(mode="digits_only"),
    )

    result = execute_plan(
        plan,
        [source_schema, compare_schema],
        {
            "F1": ("ban_ra.xlsx", source_data),
            "F2": ("ke_toan.xlsx", compare_data),
        },
    )

    assert result.stats.matched_values == 1
    output = load_workbook(BytesIO(result.files["ban_ra_da_xu_ly.xlsx"]))
    assert output["BanRa"]["A2"].fill.fill_type == "solid"
    assert output["BanRa"]["A3"].fill.fill_type is None
    output.close()


def test_oldest_person_highlights_earliest_birth_date_row():
    data = workbook_bytes(
        "NhanSu",
        ["Họ tên", "Ngày sinh"],
        [
            ["An", datetime(1990, 1, 1)],
            ["Bình", datetime(1980, 6, 15)],
            ["Chi", datetime(2000, 3, 20)],
        ],
    )
    schema = inspect_workbook(data, "nhan_su.xlsx", "F1")
    date_column = schema.sheets[0].columns[1]
    plan = OperationPlan(
        action=Action.HIGHLIGHT_EXTREME,
        source_column_id=date_column.id,
        highlight_condition="extreme",
        extreme="min",
        highlight_scope="row",
        color="FECACA",
    )

    result = execute_plan(plan, [schema], {"F1": ("nhan_su.xlsx", data)})
    output = load_workbook(BytesIO(result.files["nhan_su_da_xu_ly.xlsx"]))
    assert output["NhanSu"]["A3"].fill.fill_type == "solid"
    assert output["NhanSu"]["B3"].fill.fill_type == "solid"
    assert output["NhanSu"]["A2"].fill.fill_type is None
    output.close()


def test_invalid_tax_code_highlighting():
    data = workbook_bytes("DanhSach", ["MST"], [["0312345678"], ["12345"], ["0101234567-001"]])
    schema = inspect_workbook(data, "mst.xlsx", "F1")
    plan = OperationPlan(
        action=Action.HIGHLIGHT_INVALID_TAX_CODES,
        source_column_id=schema.sheets[0].columns[0].id,
        highlight_condition="invalid",
        color="FECACA",
    )

    result = execute_plan(plan, [schema], {"F1": ("mst.xlsx", data)})
    assert result.stats.invalid_values == 1
    output = load_workbook(BytesIO(result.files["mst_da_xu_ly.xlsx"]))
    assert output["DanhSach"]["A2"].fill.fill_type is None
    assert output["DanhSach"]["A3"].fill.fill_type == "solid"
    assert output["DanhSach"]["A4"].fill.fill_type is None
    output.close()


def test_text_normalization_handles_case_and_spaces():
    options = NormalizationOptions(
        mode="text_normalized",
        ignore_case=True,
        trim_whitespace=True,
        remove_internal_spaces=True,
    )
    assert normalize_value("  Hoa Don 01 ", options) == "hoadon01"
