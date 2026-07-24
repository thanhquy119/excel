from __future__ import annotations

from collections import Counter
from datetime import date, datetime
from io import BytesIO
from pathlib import Path
from typing import Any
from zipfile import BadZipFile, ZipFile

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

from .models import ColumnSchema, SheetSchema, WorkbookSchema

MAX_UNCOMPRESSED_BYTES = 200 * 1024 * 1024
MAX_ARCHIVE_MEMBERS = 5_000
MAX_ESTIMATED_CELLS = 2_000_000


class WorkbookValidationError(ValueError):
    pass


def validate_excel_archive(data: bytes, filename: str) -> None:
    suffix = Path(filename).suffix.lower()
    if suffix not in {".xlsx", ".xlsm"}:
        raise WorkbookValidationError("Chỉ hỗ trợ tệp .xlsx và .xlsm.")
    if not data:
        raise WorkbookValidationError("Tệp rỗng.")

    try:
        with ZipFile(BytesIO(data)) as archive:
            members = archive.infolist()
            if len(members) > MAX_ARCHIVE_MEMBERS:
                raise WorkbookValidationError("Tệp Excel có quá nhiều thành phần nội bộ.")
            total_size = sum(member.file_size for member in members)
            if total_size > MAX_UNCOMPRESSED_BYTES:
                raise WorkbookValidationError("Dung lượng giải nén của tệp vượt giới hạn an toàn 200 MB.")
            for member in members:
                normalized = member.filename.replace("\\", "/")
                if normalized.startswith("/") or "../" in normalized:
                    raise WorkbookValidationError("Tệp Excel chứa đường dẫn nội bộ không an toàn.")
    except BadZipFile as exc:
        raise WorkbookValidationError("Tệp không phải workbook Excel hợp lệ.") from exc


def _is_nonempty(value: Any) -> bool:
    return value is not None and str(value).strip() != ""


def _detect_header_row(ws, scan_rows: int = 20) -> int:
    best_row = 1
    best_score = float("-inf")
    limit = min(scan_rows, max(ws.max_row, 1))
    column_limit = min(max(ws.max_column, 1), 200)

    for row_idx in range(1, limit + 1):
        values = [ws.cell(row_idx, col_idx).value for col_idx in range(1, column_limit + 1)]
        nonempty = [value for value in values if _is_nonempty(value)]
        if not nonempty:
            continue
        text_count = sum(isinstance(value, str) for value in nonempty)
        unique_count = len({str(value).strip().casefold() for value in nonempty})
        score = len(nonempty) * 3 + text_count + unique_count - (row_idx - 1) * 0.1
        if len(nonempty) == 1:
            score -= 3
        if score > best_score:
            best_score = score
            best_row = row_idx
    return best_row


def _detect_column_type(values: list[Any]) -> str:
    populated = [value for value in values if _is_nonempty(value)]
    if not populated:
        return "empty"

    types: Counter[str] = Counter()
    for value in populated:
        if isinstance(value, bool):
            types["text"] += 1
        elif isinstance(value, (datetime, date)):
            types["date"] += 1
        elif isinstance(value, (int, float)):
            types["number"] += 1
        else:
            types["text"] += 1

    dominant, count = types.most_common(1)[0]
    return dominant if count / len(populated) >= 0.8 else "mixed"


def inspect_workbook(data: bytes, filename: str, file_id: str) -> WorkbookSchema:
    validate_excel_archive(data, filename)
    keep_vba = Path(filename).suffix.lower() == ".xlsm"
    workbook = load_workbook(
        BytesIO(data),
        read_only=True,
        data_only=False,
        keep_vba=keep_vba,
    )

    sheets: list[SheetSchema] = []
    try:
        for sheet_index, ws in enumerate(workbook.worksheets, start=1):
            estimated_cells = max(ws.max_row, 1) * max(ws.max_column, 1)
            if estimated_cells > MAX_ESTIMATED_CELLS:
                raise WorkbookValidationError(
                    f"Sheet '{ws.title}' quá lớn ({estimated_cells:,} ô). Giới hạn hiện tại là {MAX_ESTIMATED_CELLS:,} ô."
                )

            header_row = _detect_header_row(ws)
            columns: list[ColumnSchema] = []
            max_column = min(max(ws.max_column, 1), 500)
            sample_end = min(ws.max_row, header_row + 50)

            for col_idx in range(1, max_column + 1):
                raw_header = ws.cell(header_row, col_idx).value
                if not _is_nonempty(raw_header):
                    continue
                header = str(raw_header).strip()
                letter = get_column_letter(col_idx)
                sample_values = [
                    ws.cell(row_idx, col_idx).value
                    for row_idx in range(header_row + 1, sample_end + 1)
                ]
                columns.append(
                    ColumnSchema(
                        id=f"{file_id}:S{sheet_index}:{letter}",
                        letter=letter,
                        header=header,
                        detected_type=_detect_column_type(sample_values),
                    )
                )

            sheets.append(
                SheetSchema(
                    id=f"{file_id}:S{sheet_index}",
                    name=ws.title,
                    header_row=header_row,
                    max_row=ws.max_row,
                    max_column=ws.max_column,
                    columns=columns,
                )
            )
    finally:
        workbook.close()

    return WorkbookSchema(file_id=file_id, display_name=filename, sheets=sheets)


def schema_for_ai(schemas: list[WorkbookSchema]) -> list[dict[str, Any]]:
    """Return metadata only. No cell values are included."""
    return [schema.model_dump() for schema in schemas]


def build_column_index(schemas: list[WorkbookSchema]) -> dict[str, tuple[WorkbookSchema, SheetSchema, ColumnSchema]]:
    index: dict[str, tuple[WorkbookSchema, SheetSchema, ColumnSchema]] = {}
    for workbook in schemas:
        for sheet in workbook.sheets:
            for column in sheet.columns:
                index[column.id] = (workbook, sheet, column)
    return index
