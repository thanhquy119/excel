from __future__ import annotations

import re
import unicodedata
from collections import Counter
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from io import BytesIO
from pathlib import Path
from typing import Any, Callable

from openpyxl import load_workbook
from openpyxl.styles import PatternFill

from .models import Action, ExecutionResult, ExecutionStats, NormalizationOptions, OperationPlan, WorkbookSchema
from .schema import build_column_index


class OperationError(ValueError):
    pass


def normalize_value(value: Any, options: NormalizationOptions) -> Any:
    if value is None:
        return None

    if options.mode == "exact":
        if isinstance(value, str) and options.trim_whitespace:
            value = value.strip()
        return (type(value).__name__, value)

    text = unicodedata.normalize("NFKC", str(value))
    if options.trim_whitespace:
        text = text.strip()
    if options.remove_internal_spaces:
        text = re.sub(r"\s+", "", text)
    if options.ignore_case:
        text = text.casefold()

    if text == "":
        return None
    if options.mode == "digits_only":
        digits = re.sub(r"\D", "", text)
        return digits or None
    if options.mode == "number":
        return _parse_number(value)
    return text


def _parse_number(value: Any) -> Decimal | None:
    if value is None or isinstance(value, bool):
        return None
    if isinstance(value, Decimal):
        return value
    if isinstance(value, (int, float)):
        return Decimal(str(value))

    text = unicodedata.normalize("NFKC", str(value)).strip().replace(" ", "")
    if not text:
        return None

    negative = text.startswith("(") and text.endswith(")")
    if negative:
        text = text[1:-1]
    text = re.sub(r"[^0-9,\.\-+]", "", text)

    if "," in text and "." in text:
        if text.rfind(",") > text.rfind("."):
            text = text.replace(".", "").replace(",", ".")
        else:
            text = text.replace(",", "")
    elif "," in text:
        parts = text.split(",")
        if len(parts) == 2 and 0 < len(parts[1]) <= 2:
            text = ".".join(parts)
        else:
            text = "".join(parts)
    elif text.count(".") > 1:
        parts = text.split(".")
        if len(parts[-1]) <= 2:
            text = "".join(parts[:-1]) + "." + parts[-1]
        else:
            text = "".join(parts)

    try:
        number = Decimal(text)
        return -number if negative else number
    except (InvalidOperation, ValueError):
        return None


def _parse_date(value: Any) -> date | None:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if value is None:
        return None

    text = unicodedata.normalize("NFKC", str(value)).strip()
    if not text:
        return None

    formats = (
        "%Y-%m-%d",
        "%d/%m/%Y",
        "%d-%m-%Y",
        "%Y/%m/%d",
        "%d.%m.%Y",
        "%m/%d/%Y",
    )
    for fmt in formats:
        try:
            return datetime.strptime(text, fmt).date()
        except ValueError:
            continue
    return None


def _fill(color: str) -> PatternFill:
    rgb = color.upper()
    if len(rgb) == 6:
        rgb = f"FF{rgb}"
    return PatternFill(fill_type="solid", fgColor=rgb)


def _load_editable(filename: str, data: bytes):
    keep_vba = Path(filename).suffix.lower() == ".xlsm"
    return load_workbook(BytesIO(data), data_only=False, keep_vba=keep_vba)


def _output_name(filename: str) -> str:
    path = Path(filename)
    return f"{path.stem}_da_xu_ly{path.suffix.lower()}"


def _iter_cells(ws, column_letter: str, header_row: int):
    for row_idx in range(header_row + 1, ws.max_row + 1):
        yield ws[f"{column_letter}{row_idx}"]


def _save_changed(workbooks: dict[str, Any], filenames: dict[str, str], changed_ids: set[str]) -> dict[str, bytes]:
    files: dict[str, bytes] = {}
    for file_id in changed_ids:
        buffer = BytesIO()
        workbooks[file_id].save(buffer)
        files[_output_name(filenames[file_id])] = buffer.getvalue()
    return files


def _validate_plan_columns(plan: OperationPlan, schemas: list[WorkbookSchema]):
    index = build_column_index(schemas)
    if plan.source_column_id not in index:
        raise OperationError("Cột nguồn không tồn tại trong workbook hiện tại.")
    if plan.comparison_column_id and plan.comparison_column_id not in index:
        raise OperationError("Cột đối chiếu không tồn tại trong workbook hiện tại.")
    return index


def _apply_fill(ws, cell, cell_fill: PatternFill, scope: str) -> int:
    if scope == "row":
        count = 0
        for column_index in range(1, ws.max_column + 1):
            ws.cell(cell.row, column_index).fill = cell_fill
            count += 1
        return count
    cell.fill = cell_fill
    return 1


def execute_plan(
    plan: OperationPlan,
    schemas: list[WorkbookSchema],
    uploaded_files: dict[str, tuple[str, bytes]],
) -> ExecutionResult:
    if plan.needs_confirmation:
        raise OperationError(plan.question or "Yêu cầu chưa đủ rõ để xử lý an toàn.")

    column_index = _validate_plan_columns(plan, schemas)
    filenames = {file_id: item[0] for file_id, item in uploaded_files.items()}
    workbooks = {
        file_id: _load_editable(filename, data)
        for file_id, (filename, data) in uploaded_files.items()
    }

    try:
        handlers: dict[Action, Callable[..., ExecutionResult]] = {
            Action.COMPARE_COLUMNS: _compare_columns,
            Action.HIGHLIGHT_DUPLICATES: _highlight_duplicates,
            Action.FIND_BLANKS: _find_blanks,
            Action.HIGHLIGHT_INVALID_TAX_CODES: _highlight_invalid_tax_codes,
            Action.HIGHLIGHT_NUMBER_CONDITION: _highlight_number_condition,
            Action.HIGHLIGHT_DATE_CONDITION: _highlight_date_condition,
            Action.HIGHLIGHT_TEXT_CONDITION: _highlight_text_condition,
            Action.HIGHLIGHT_EXTREME: _highlight_extreme,
        }
        handler = handlers.get(plan.action)
        if not handler:
            raise OperationError("Hành động chưa được hỗ trợ.")
        return handler(plan, column_index, workbooks, filenames)
    finally:
        for workbook in workbooks.values():
            workbook.close()


def _source_context(plan, column_index, workbooks):
    source_book, source_sheet, source_column = column_index[plan.source_column_id]
    ws = workbooks[source_book.file_id][source_sheet.name]
    cells = list(_iter_cells(ws, source_column.letter, source_sheet.header_row))
    return source_book, source_sheet, source_column, ws, cells


def _compare_columns(plan, column_index, workbooks, filenames) -> ExecutionResult:
    source_book, source_sheet, source_column = column_index[plan.source_column_id]
    compare_book, compare_sheet, compare_column = column_index[plan.comparison_column_id]
    source_ws = workbooks[source_book.file_id][source_sheet.name]
    compare_ws = workbooks[compare_book.file_id][compare_sheet.name]

    source_cells = list(_iter_cells(source_ws, source_column.letter, source_sheet.header_row))
    compare_cells = list(_iter_cells(compare_ws, compare_column.letter, compare_sheet.header_row))
    source_normalized = [normalize_value(cell.value, plan.normalization) for cell in source_cells]
    compare_normalized = [normalize_value(cell.value, plan.normalization) for cell in compare_cells]
    source_set = {value for value in source_normalized if value is not None}
    compare_set = {value for value in compare_normalized if value is not None}

    matched = sum(value is not None and value in compare_set for value in source_normalized)
    unmatched = sum(value is not None and value not in compare_set for value in source_normalized)
    highlighted = 0
    changed_ids: set[str] = set()
    cell_fill = _fill(plan.color)

    if plan.highlight_target in {"source", "both"}:
        for cell, value in zip(source_cells, source_normalized):
            is_match = value is not None and value in compare_set
            should_highlight = is_match if plan.highlight_condition == "matched" else value is not None and not is_match
            if should_highlight:
                highlighted += _apply_fill(source_ws, cell, cell_fill, plan.highlight_scope)
        changed_ids.add(source_book.file_id)

    if plan.highlight_target in {"comparison", "both"}:
        for cell, value in zip(compare_cells, compare_normalized):
            is_match = value is not None and value in source_set
            should_highlight = is_match if plan.highlight_condition == "matched" else value is not None and not is_match
            if should_highlight:
                highlighted += _apply_fill(compare_ws, cell, cell_fill, plan.highlight_scope)
        changed_ids.add(compare_book.file_id)

    files = _save_changed(workbooks, filenames, changed_ids)
    return ExecutionResult(
        files=files,
        stats=ExecutionStats(
            scanned_rows=len(source_cells) + len(compare_cells),
            highlighted_cells=highlighted,
            matched_values=matched,
            unmatched_values=unmatched,
        ),
        message="Đã xử lý đối chiếu hai cột.",
    )


def _highlight_duplicates(plan, column_index, workbooks, filenames) -> ExecutionResult:
    source_book, _, _, ws, cells = _source_context(plan, column_index, workbooks)
    normalized = [normalize_value(cell.value, plan.normalization) for cell in cells]
    counts = Counter(value for value in normalized if value is not None)
    duplicate_keys = {value for value, count in counts.items() if count > 1}
    cell_fill = _fill(plan.color)
    highlighted = 0

    for cell, value in zip(cells, normalized):
        if value in duplicate_keys:
            highlighted += _apply_fill(ws, cell, cell_fill, plan.highlight_scope)

    files = _save_changed(workbooks, filenames, {source_book.file_id})
    return ExecutionResult(
        files=files,
        stats=ExecutionStats(
            scanned_rows=len(cells),
            highlighted_cells=highlighted,
            duplicate_values=len(duplicate_keys),
        ),
        message="Đã đánh dấu dữ liệu trùng.",
    )


def _find_blanks(plan, column_index, workbooks, filenames) -> ExecutionResult:
    source_book, _, _, ws, cells = _source_context(plan, column_index, workbooks)
    cell_fill = _fill(plan.color)
    blanks = 0
    highlighted = 0

    for cell in cells:
        if cell.value is None or str(cell.value).strip() == "":
            highlighted += _apply_fill(ws, cell, cell_fill, plan.highlight_scope)
            blanks += 1

    files = _save_changed(workbooks, filenames, {source_book.file_id})
    return ExecutionResult(
        files=files,
        stats=ExecutionStats(scanned_rows=len(cells), highlighted_cells=highlighted, blank_cells=blanks),
        message="Đã đánh dấu ô trống.",
    )


def _highlight_invalid_tax_codes(plan, column_index, workbooks, filenames) -> ExecutionResult:
    source_book, _, _, ws, cells = _source_context(plan, column_index, workbooks)
    allowed_lengths = set(plan.tax_code_lengths)
    cell_fill = _fill(plan.color)
    invalid = 0
    highlighted = 0

    for cell in cells:
        raw = "" if cell.value is None else str(cell.value).strip()
        if not raw:
            continue
        digits = re.sub(r"\D", "", raw)
        valid = len(digits) in allowed_lengths and len(re.sub(r"[0-9\s.\-]", "", raw)) == 0
        if not valid:
            highlighted += _apply_fill(ws, cell, cell_fill, plan.highlight_scope)
            invalid += 1

    files = _save_changed(workbooks, filenames, {source_book.file_id})
    return ExecutionResult(
        files=files,
        stats=ExecutionStats(scanned_rows=len(cells), highlighted_cells=highlighted, invalid_values=invalid),
        message="Đã đánh dấu mã số thuế không hợp lệ về định dạng.",
    )


def _number_matches(value: Decimal, plan: OperationPlan) -> bool:
    first = Decimal(str(plan.number_value))
    second = Decimal(str(plan.second_number_value)) if plan.second_number_value is not None else None
    operator = plan.number_operator
    if operator == "gt":
        return value > first
    if operator == "gte":
        return value >= first
    if operator == "lt":
        return value < first
    if operator == "lte":
        return value <= first
    if operator == "eq":
        return value == first
    if operator == "ne":
        return value != first
    if operator == "between" and second is not None:
        lower, upper = sorted((first, second))
        return lower <= value <= upper
    return False


def _highlight_number_condition(plan, column_index, workbooks, filenames) -> ExecutionResult:
    source_book, _, _, ws, cells = _source_context(plan, column_index, workbooks)
    cell_fill = _fill(plan.color)
    matches = 0
    highlighted = 0

    for cell in cells:
        number = _parse_number(cell.value)
        if number is not None and _number_matches(number, plan):
            highlighted += _apply_fill(ws, cell, cell_fill, plan.highlight_scope)
            matches += 1

    files = _save_changed(workbooks, filenames, {source_book.file_id})
    return ExecutionResult(
        files=files,
        stats=ExecutionStats(scanned_rows=len(cells), highlighted_cells=highlighted, conditional_values=matches),
        message="Đã đánh dấu các giá trị thỏa điều kiện số.",
    )


def _date_matches(value: date, plan: OperationPlan) -> bool:
    first = _parse_date(plan.date_value)
    second = _parse_date(plan.second_date_value) if plan.second_date_value else None
    if first is None:
        raise OperationError("Ngày trong yêu cầu không hợp lệ.")
    operator = plan.date_operator
    if operator == "before":
        return value < first
    if operator == "after":
        return value > first
    if operator == "on":
        return value == first
    if operator == "between" and second is not None:
        lower, upper = sorted((first, second))
        return lower <= value <= upper
    return False


def _highlight_date_condition(plan, column_index, workbooks, filenames) -> ExecutionResult:
    source_book, _, _, ws, cells = _source_context(plan, column_index, workbooks)
    cell_fill = _fill(plan.color)
    matches = 0
    highlighted = 0

    for cell in cells:
        parsed = _parse_date(cell.value)
        if parsed is not None and _date_matches(parsed, plan):
            highlighted += _apply_fill(ws, cell, cell_fill, plan.highlight_scope)
            matches += 1

    files = _save_changed(workbooks, filenames, {source_book.file_id})
    return ExecutionResult(
        files=files,
        stats=ExecutionStats(scanned_rows=len(cells), highlighted_cells=highlighted, conditional_values=matches),
        message="Đã đánh dấu các ngày thỏa điều kiện.",
    )


def _text_matches(value: Any, plan: OperationPlan) -> bool:
    if value is None:
        return False
    text = unicodedata.normalize("NFKC", str(value))
    expected = unicodedata.normalize("NFKC", plan.text_value or "")
    if plan.normalization.trim_whitespace:
        text = text.strip()
        expected = expected.strip()
    if plan.normalization.ignore_case:
        text = text.casefold()
        expected = expected.casefold()

    operator = plan.text_operator
    if operator == "contains":
        return expected in text
    if operator == "not_contains":
        return expected not in text
    if operator == "starts_with":
        return text.startswith(expected)
    if operator == "ends_with":
        return text.endswith(expected)
    if operator == "equals":
        return text == expected
    return False


def _highlight_text_condition(plan, column_index, workbooks, filenames) -> ExecutionResult:
    source_book, _, _, ws, cells = _source_context(plan, column_index, workbooks)
    cell_fill = _fill(plan.color)
    matches = 0
    highlighted = 0

    for cell in cells:
        if _text_matches(cell.value, plan):
            highlighted += _apply_fill(ws, cell, cell_fill, plan.highlight_scope)
            matches += 1

    files = _save_changed(workbooks, filenames, {source_book.file_id})
    return ExecutionResult(
        files=files,
        stats=ExecutionStats(scanned_rows=len(cells), highlighted_cells=highlighted, conditional_values=matches),
        message="Đã đánh dấu các ô thỏa điều kiện văn bản.",
    )


def _highlight_extreme(plan, column_index, workbooks, filenames) -> ExecutionResult:
    source_book, source_sheet, source_column, ws, cells = _source_context(plan, column_index, workbooks)

    if source_column.detected_type == "date":
        converted = [_parse_date(cell.value) for cell in cells]
    elif source_column.detected_type == "number":
        converted = [_parse_number(cell.value) for cell in cells]
    else:
        header = source_column.header.casefold()
        if any(word in header for word in ("ngày", "date", "sinh")):
            converted = [_parse_date(cell.value) for cell in cells]
        else:
            converted = [_parse_number(cell.value) for cell in cells]

    populated = [value for value in converted if value is not None]
    if not populated:
        raise OperationError(f"Không tìm thấy giá trị ngày hoặc số hợp lệ trong sheet '{source_sheet.name}'.")

    extreme_value = min(populated) if plan.extreme == "min" else max(populated)
    cell_fill = _fill(plan.color)
    matches = 0
    highlighted = 0

    for cell, value in zip(cells, converted):
        if value == extreme_value:
            highlighted += _apply_fill(ws, cell, cell_fill, plan.highlight_scope)
            matches += 1

    files = _save_changed(workbooks, filenames, {source_book.file_id})
    return ExecutionResult(
        files=files,
        stats=ExecutionStats(scanned_rows=len(cells), highlighted_cells=highlighted, extreme_values=matches),
        message="Đã đánh dấu giá trị lớn nhất hoặc nhỏ nhất theo yêu cầu.",
    )
